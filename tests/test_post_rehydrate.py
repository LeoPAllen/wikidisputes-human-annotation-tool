from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path

from openpyxl import load_workbook
import pandas as pd

from wikidisputes_ui.codebook import load_codebook, schema_id
from wikidisputes_ui.export import build_export
from wikidisputes_ui.ingest import Dataset
from wikidisputes_ui.models import normalize_and_validate
from wikidisputes_ui.source_migration import reconcile_annotation_keys
from wikidisputes_ui.storage import Storage


def dataset(rows: list[dict[str, object]]) -> Dataset:
    frame = pd.DataFrame(rows)
    frame["utterance_role"] = "utterance"
    frame["dispute_sequence"] = 1
    frame["timestamp"] = "2026-01-01"
    frame["utterance_type"] = "talk"
    frame["article_title"] = "Article"
    frame["_source_row"] = range(2, len(frame) + 2)
    return Dataset(frame)


def row(uid: str, order: int, text: str, *, original: str | None = None) -> dict[str, object]:
    return {
        "dispute_id": "D1",
        "utterance_order": order,
        "utterance_id": uid,
        "original_utterance_id": original,
        "speaker_id": "A",
        "reply_to_utterance_id": None,
        "utterance_text": text,
    }


def save(storage: Storage, uid: str, *, saved: str = "2026-01-01T00:00:00Z", marker: int = 0) -> None:
    storage.set_active_coder("coder_1")
    storage.save_utterance(
        coder="coder_1",
        utterance_id=uid,
        dispute_id="D1",
        payload={
            "KS_present": marker,
            "KS_explicit_reasoning": None,
            "KS_grounding": None,
            "KS_restaking": None,
            "KS_bounding": None,
            "KI_present": 0,
            "C_off_topic_shift": 0,
            "C_interpersonal_attack_or_disrespect": 0,
            "C_formal_governance_action": 0,
            "malformed_utterance": False,
        },
        answered_fields={"KS_present", "malformed_utterance"},
        submit=True,
        schema_version="old",
        schema_hash="old-hash",
        opened_at=saved,
        elapsed_wall_seconds=1,
    )
    with storage.connect() as db:
        db.execute(
            "UPDATE utterance_annotations SET saved_at=? WHERE coder_id='coder_1' AND utterance_id=?",
            (saved, uid),
        )


def test_annotation_survives_reordered_text_changed_gold(tmp_path: Path) -> None:
    storage = Storage(tmp_path / "annotations.sqlite3")
    old = dataset([row("u1", 1, "old", original="stable-1"), row("u2", 2, "two", original="stable-2")])
    save(storage, "u1")
    reconcile_annotation_keys(storage, old)

    replacement = dataset([row("u2", 1, "two changed", original="stable-2"), row("u1", 2, "new", original="stable-1")])
    reconcile_annotation_keys(storage, replacement)
    assert storage.current_utterance("coder_1", "stable-1")["status"] == "submitted"
    assert replacement.annotatable_rows["utterance_order"].tolist() == [1, 2]


def test_changed_current_id_same_stable_key_survives(tmp_path: Path) -> None:
    storage = Storage(tmp_path / "annotations.sqlite3")
    save(storage, "old-current")
    reconcile_annotation_keys(storage, dataset([row("old-current", 1, "old", original="stable")]))
    reconcile_annotation_keys(storage, dataset([row("new-current", 1, "new", original="stable")]))
    assert storage.current_utterance("coder_1", "stable") is not None
    assert storage.current_utterance("coder_1", "old-current") is None


def test_add_remove_and_collapse_preserve_work(tmp_path: Path) -> None:
    storage = Storage(tmp_path / "annotations.sqlite3")
    save(storage, "old-a", saved="2026-01-01T00:00:00Z", marker=0)
    save(storage, "old-b", saved="2026-01-02T00:00:00Z", marker=1)
    save(storage, "removed", marker=0)
    before_events = storage.rows("utterance_annotation_events", "coder_1")
    old = dataset(
        [
            row("old-a", 1, "a", original="collapsed"),
            row("old-b", 2, "b", original="collapsed"),
            row("removed", 3, "gone"),
        ]
    )
    reconcile_annotation_keys(storage, old)
    current = dataset([row("current", 1, "combined", original="collapsed"), row("new", 2, "new")])
    reconcile_annotation_keys(storage, current)

    collapsed = storage.current_utterance("coder_1", "collapsed")
    assert collapsed["payload"]["KS_present"] == 1
    assert storage.current_utterance("coder_1", "new") is None
    assert storage.current_utterance("coder_1", "removed") is not None
    assert storage.rows("utterance_annotation_events", "coder_1") == before_events
    exported = pd.read_excel(BytesIO(build_export(storage, current, "coder_1", "new", "new-hash", ("KS_present",), ())))
    assert list(exported["utterance_id"]) == ["current"]


def test_changed_codebook_bytes_register_normally(tmp_path: Path, synthetic_project) -> None:
    original = synthetic_project.codebook_path
    changed = tmp_path / "changed-codebook.xlsx"
    changed.write_bytes(original.read_bytes())
    workbook = load_workbook(changed)
    sheet = workbook["Core_Schema"]
    sheet.cell(row=2, column=4).value = f"{sheet.cell(row=2, column=4).value} "
    workbook.save(changed)
    first, second = load_codebook(original), load_codebook(changed)
    assert first.file_hash != second.file_hash
    storage = Storage(tmp_path / "annotations.sqlite3")
    storage.register_schema(schema_id(first.file_hash), first.file_hash, first.source_filename)
    storage.register_schema(schema_id(second.file_hash), second.file_hash, second.source_filename)


def test_malformed_flag_persists_exports_and_allows_missing_constructs(tmp_path: Path) -> None:
    values = {"malformed_utterance": True, "coder_confidence": 3, "review_flag": 0}
    result = normalize_and_validate(values, set(values))
    assert result.valid
    storage = Storage(tmp_path / "annotations.sqlite3")
    storage.set_active_coder("coder_1")
    storage.save_utterance(
        coder="coder_1",
        utterance_id="stable",
        dispute_id="D1",
        payload=result.payload,
        answered_fields=set(values),
        submit=True,
        schema_version="schema",
        schema_hash="hash",
        opened_at="2026-01-01T00:00:00Z",
        elapsed_wall_seconds=1,
    )
    current = storage.current_utterance("coder_1", "stable")
    event = storage.rows("utterance_annotation_events", "coder_1")[0]
    assert current["payload"]["malformed_utterance"] is True
    assert json.loads(event["payload_json"])["malformed_utterance"] is True
    source = dataset([row("current-id", 1, "bad", original="stable")])
    exported = pd.read_excel(BytesIO(build_export(storage, source, "coder_1", "schema", "hash", (), ())))
    assert bool(exported.loc[0, "malformed_utterance"])
